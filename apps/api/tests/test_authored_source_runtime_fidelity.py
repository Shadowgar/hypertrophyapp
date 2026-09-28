"""Independent licensed-source qualification; never republish workbook rows.

The OpenPyXL reader does not use the importer or committed artifacts to decide
which rows exist. Missing locally authorized workbooks are an explicit skip.
Synthetic typed-target/receipt coverage is in test_authored_typed_execution.py.
"""
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re

import pytest
from openpyxl import load_workbook

from app.program_loader import load_program_template
from core_engine.scheduler import generate_week_plan, AUTHORITATIVE_AUTHORED_PASSTHROUGH_KEY
from core_engine.decision_workout_session import build_workout_today_payload
from importers.structured_program_builder import collect_structured_phases_from_workbook
from importers.authored_lineage import artifact_sha256, compiler_sha256

ROOT = Path(__file__).resolve().parents[3]
# Fixed source table columns, independent of importer's header matching.
COLUMNS = {"exercise": 2, "last_set_intensity_technique": 3, "warm_up_sets": 4,
    "working_sets": 5, "reps": 6, "tracking_set_1": 7, "tracking_set_2": 8,
    "tracking_set_3": 9, "tracking_set_4": 10, "early_set_rpe": 11,
    "last_set_rpe": 12, "rest": 13, "substitution_option_1": 14,
    "substitution_option_2": 15, "notes": 16}


def cell(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return f"{value.month}-{value.day}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def comparable(value):
    normalized = re.sub(r"\s+", " ", value or "").strip()
    return normalized[:-2] if re.fullmatch(r"\d+\.0", normalized) else normalized


def independent_rows(path):
    workbook = load_workbook(path, data_only=True)
    sheet = workbook["Full Body"]
    rows = []
    week = None
    day = None
    order = 0
    for row_index in range(1, sheet.max_row + 1):
        first = cell(sheet.cell(row_index, 1).value)
        name = cell(sheet.cell(row_index, 2).value)
        if first and re.fullmatch(r"Week \d+", first):
            week = int(first.split()[1]); day = None; order = 0
            continue
        if week is None or not name or name == "Exercise" or first == "Mandatory Rest Day" or (first or "").startswith(("BLOCK", "SEMI-DELOAD")):
            continue
        if first:
            day = first; order = 0
        if day is None:
            continue
        order += 1
        rows.append({"week": week, "day": day, "order": order, "row": row_index,
            "raw": {field: cell(sheet.cell(row_index, column).value) for field, column in COLUMNS.items()}})
    workbook.close()
    return rows


@pytest.mark.parametrize("phase,expected_count", [(1, 315), (2, 310)])
def test_source_import_canonical_runtime_today_parity(phase, expected_count):
    source = ROOT / "reference" / f"Pure Bodybuilding Phase {phase} - Full Body Sheet.xlsx"
    if not source.exists():
        pytest.skip("Locally authorized licensed workbook unavailable; source parity not qualified in this environment")
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    source_rows = independent_rows(source)
    assert len(source_rows) == expected_count
    program_id = f"pure_bodybuilding_phase_{phase}_full_body"
    imported, _, _ = collect_structured_phases_from_workbook(source, sheet_name="Full Body")
    import_rows = [exercise for p in imported for w in p["weeks"] for s in w["sessions"] for exercise in s["exercises"]]
    canonical = json.loads((ROOT / "programs/gold" / f"{program_id}.json").read_text())
    canonical_rows = [slot for p in canonical["phases"] for w in p["weeks"] for d in w["days"] for slot in d["slots"]]
    template = load_program_template(program_id)
    runtime_rows = [exercise for w in template["authored_weeks"] for s in w["sessions"] for exercise in s["exercises"]]
    template[AUTHORITATIVE_AUTHORED_PASSTHROUGH_KEY] = True
    execution_rows = []
    for week in range(10):
        payload = generate_week_plan({"name": "Synthetic fidelity"}, 5, "full_body", template, [], "maintenance", prior_generated_weeks=week)
        for session in payload["sessions"]:
            today = build_workout_today_payload(selected_session=session, mesocycle=None, deload=None,
                completed_sets_by_exercise={}, live_recommendations_by_exercise={}, resume_selected=False, daily_quote={})
            execution_rows.extend(today["exercises"])
    stages = {"import": import_rows, "canonical": canonical_rows, "runtime": runtime_rows, "execution": execution_rows}
    assert all(len(rows) == expected_count for rows in stages.values()), {stage: len(rows) for stage, rows in stages.items()}
    assert canonical["source_provenance"]["source_sha256"] == before
    assert canonical["source_provenance"]["importer_sha256"] == compiler_sha256()
    assert canonical["source_provenance"]["artifact_sha256"] == artifact_sha256(canonical)
    lineage_ids = set()
    amrap_weeks = []
    positional_sets = 0
    for source_index, row in enumerate(source_rows):
        location = (phase, row["week"], row["day"], row["order"])
        # A source count may explicitly qualify a unilateral set as "per leg".
        count_match = re.fullmatch(r"(\d+)(?: per (?:leg|arm|side))?", row["raw"]["working_sets"])
        assert count_match is not None, location
        count = int(count_match[1])
        rep_parts = re.split(r"\s*[,;/]\s*", row["raw"]["reps"] or "")
        expected_reps = rep_parts if len(rep_parts) == count else [row["raw"]["reps"]] * count
        positional_sets += int(len(rep_parts) == count and len(set(rep_parts)) > 1)
        for stage, rows in stages.items():
            slot = rows[source_index]
            prescribed = slot["authored_prescription"]
            for field in COLUMNS:
                assert comparable(prescribed["raw"][field]) == comparable(row["raw"][field]), (location, stage, field)
            assert len(prescribed["sets"]) == count, (location, stage, "working sets")
            for set_index, item in enumerate(prescribed["sets"]):
                assert comparable(item["rep_target"]["raw"]) == comparable(expected_reps[set_index]), (location, stage, "rep target", set_index)
                expected_effort = row["raw"]["last_set_rpe"] if set_index == count - 1 else row["raw"]["early_set_rpe"]
                assert comparable(item["effort_target"]["raw"]) == comparable(expected_effort), (location, stage, "effort", set_index)
                assert comparable(item["rest"]) == comparable(row["raw"]["rest"]), (location, stage, "rest")
                expected_technique = row["raw"]["last_set_intensity_technique"] if set_index == count - 1 else None
                assert comparable(item["intensity_technique"]) == comparable(expected_technique), (location, stage, "intensity")
        assert canonical_rows[source_index]["warmup_prescription"] == [], location
        assert execution_rows[source_index]["warmups"] == [], location
        lineage = runtime_rows[source_index]["source_lineage"]
        assert lineage["source_week"] == row["week"] and lineage["source_slot"] == row["order"]
        assert lineage["source_program_id"] == program_id and lineage["source_row"] == row["row"]
        assert lineage["source_sha256"] == before and lineage["artifact_sha256"] == canonical["source_provenance"]["artifact_sha256"]
        assert lineage["source_slot_id"] not in lineage_ids
        lineage_ids.add(lineage["source_slot_id"])
        for item in runtime_rows[source_index]["authored_prescription"]["sets"]:
            assert item["source_set_id"] == f"{lineage['source_slot_id']}:set{item['set_index']}"
        assert execution_rows[source_index]["source_lineage"] == lineage
        if row["raw"]["reps"] == "AMRAP":
            amrap_weeks.append(row["week"])
            assert count == 2
            for stage, rows in stages.items():
                assert all(item["rep_target"]["kind"] == "amrap" for item in rows[source_index]["authored_prescription"]["sets"]), (stage, location)
                if stage in ("import", "runtime", "execution"):
                    assert rows[source_index]["rep_range"] is None
    if phase == 1:
        assert positional_sets >= 10
    else:
        assert amrap_weeks == [6, 7, 8, 9, 10]
    assert before == hashlib.sha256(source.read_bytes()).hexdigest()
