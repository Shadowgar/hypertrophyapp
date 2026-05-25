from __future__ import annotations

from datetime import datetime
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[3]
PHASE1_WORKBOOK = REPO_ROOT / "reference" / "Pure Bodybuilding Phase 1 - Full Body Sheet.xlsx"
PHASE1_ONBOARDING = REPO_ROOT / "programs" / "gold" / "pure_bodybuilding_phase_1_full_body.onboarding.json"
PHASE1_TEMPLATE = REPO_ROOT / "programs" / "gold" / "pure_bodybuilding_phase_1_full_body.json"

SLOT_FIELDS = [
    "exercise",
    "last_set_intensity_technique",
    "warm_up_sets",
    "working_sets",
    "reps",
    "early_set_rpe",
    "last_set_rpe",
    "rest",
    "substitution_option_1",
    "substitution_option_2",
    "notes",
]


def _cell_to_source_string(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        # Some Excel cells such as "2-3" are stored as dates by the source workbook.
        return f"{value.month}-{value.day}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize(value: Any) -> str:
    text = _cell_to_source_string(value)
    if text is None:
        return ""
    text = (
        text.replace("°", "deg")
        .replace("–", "-")
        .replace("—", "-")
        .replace("“", '"')
        .replace("”", '"')
        .replace("’", "'")
    )
    text = re.sub(r"\s+", " ", text).strip()
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def _parse_workbook_slots() -> list[dict[str, Any]]:
    workbook = load_workbook(PHASE1_WORKBOOK, data_only=True)
    sheet = workbook["Full Body"]
    parsed: list[dict[str, Any]] = []
    current_week: int | None = None
    current_day: str | None = None
    current_order = 0

    for row_index in range(1, sheet.max_row + 1):
        first = _cell_to_source_string(sheet.cell(row_index, 1).value)
        exercise = _cell_to_source_string(sheet.cell(row_index, 2).value)

        if first and re.fullmatch(r"Week \d+", first):
            current_week = int(first.split()[1])
            current_day = None
            current_order = 0
            continue

        if (
            current_week is None
            or not exercise
            or exercise == "Exercise"
            or first == "Mandatory Rest Day"
            or (first or "").startswith("BLOCK")
            or (first or "").startswith("SEMI-DELOAD")
        ):
            continue

        if first:
            current_day = first
            current_order = 0
        if current_day is None:
            continue

        current_order += 1
        parsed.append(
            {
                "week": current_week,
                "day": current_day,
                "order": current_order,
                "exercise": exercise,
                "last_set_intensity_technique": _cell_to_source_string(sheet.cell(row_index, 3).value),
                "warm_up_sets": _cell_to_source_string(sheet.cell(row_index, 4).value),
                "working_sets": _cell_to_source_string(sheet.cell(row_index, 5).value),
                "reps": _cell_to_source_string(sheet.cell(row_index, 6).value),
                "early_set_rpe": _cell_to_source_string(sheet.cell(row_index, 11).value),
                "last_set_rpe": _cell_to_source_string(sheet.cell(row_index, 12).value),
                "rest": _cell_to_source_string(sheet.cell(row_index, 13).value),
                "substitution_option_1": _cell_to_source_string(sheet.cell(row_index, 14).value),
                "substitution_option_2": _cell_to_source_string(sheet.cell(row_index, 15).value),
                "notes": _cell_to_source_string(sheet.cell(row_index, 16).value),
            }
        )

    return parsed


def _onboarding_slots() -> list[dict[str, Any]]:
    payload = json.loads(PHASE1_ONBOARDING.read_text(encoding="utf-8"))
    parsed: list[dict[str, Any]] = []
    for week_index, week in enumerate(payload["blueprint"]["week_templates"], start=1):
        for day in week["days"]:
            for slot in day["slots"]:
                parsed.append(
                    {
                        "week": week_index,
                        "day": day["day_name"],
                        "order": slot["order_index"],
                        **{field: slot.get(field) for field in SLOT_FIELDS},
                    }
                )
    return parsed


def test_phase1_full_body_source_paths_point_to_real_companion_workbook() -> None:
    assert PHASE1_WORKBOOK.exists()

    template = json.loads(PHASE1_TEMPLATE.read_text(encoding="utf-8"))
    onboarding = json.loads(PHASE1_ONBOARDING.read_text(encoding="utf-8"))

    expected_path = str(PHASE1_WORKBOOK)
    assert template["source_workbook"] == expected_path
    assert onboarding["blueprint"]["source_workbook"] == expected_path
    assert Path(onboarding["source_pdf"]).name == "The_Pure_Bodybuilding_Program - Phase 1 - Full_Body.pdf"


def test_phase1_full_body_onboarding_matches_companion_workbook_table_exactly() -> None:
    workbook_slots = _parse_workbook_slots()
    onboarding_slots = _onboarding_slots()

    assert len(workbook_slots) == 315
    assert len(onboarding_slots) == 315

    for source_slot, app_slot in zip(workbook_slots, onboarding_slots, strict=True):
        location = f"week {source_slot['week']} {source_slot['day']} slot {source_slot['order']}"
        assert (app_slot["week"], app_slot["day"], app_slot["order"]) == (
            source_slot["week"],
            source_slot["day"],
            source_slot["order"],
        ), location
        for field in SLOT_FIELDS:
            assert _normalize(app_slot.get(field)) == _normalize(source_slot.get(field)), (
                location,
                field,
                source_slot.get(field),
                app_slot.get(field),
            )

